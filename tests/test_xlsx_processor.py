"""Минимальные тесты на парсер xlsx.

Запуск: `python -m unittest tests.test_xlsx_processor` из корня проекта.
"""
from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

from src.xlsx_processor import (
    COMBINED_FILL,
    HIGHLIGHT_FILL,
    WARNING_FILL,
    Correction,
    WarningCell,
    parse,
    write_result,
)


ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "Сизикова Кристина Андреевна (2).xlsx"


class TestParse(unittest.TestCase):
    def test_real_file_has_expected_entries(self) -> None:
        if not SAMPLE.exists():
            self.skipTest(f"Sample file not found: {SAMPLE}")

        _wb, entries, header_row = parse(SAMPLE)

        self.assertEqual(len(entries), 151, "ожидается 151 строка работ")
        self.assertEqual(header_row, 5, "строка заголовков = 5")
        for e in entries:
            self.assertTrue(e.contractor, f"row {e.row_idx} без контрагента")
            self.assertTrue(e.content, f"row {e.row_idx} без содержания")
            self.assertTrue(e.date, f"row {e.row_idx} без даты")
            self.assertEqual(e.content_col, 4, "колонка «Содержание» = D")

    def test_first_entry_fields(self) -> None:
        if not SAMPLE.exists():
            self.skipTest(f"Sample file not found: {SAMPLE}")

        _wb, entries, _header = parse(SAMPLE)
        first = entries[0]
        self.assertEqual(first.contractor, "Aqua Trade ИП")
        self.assertEqual(first.date, "14.05.2026")
        self.assertEqual(first.time, "17:05:00")
        self.assertTrue(first.content.startswith("Архив ИБ БК"))

    def test_contractor_keeps_for_multiple_dates(self) -> None:
        if not SAMPLE.exists():
            self.skipTest(f"Sample file not found: {SAMPLE}")

        _wb, entries, _header = parse(SAMPLE)
        expert_rows = [e for e in entries if e.contractor == "EXPERT-SK"]
        self.assertGreaterEqual(
            len(expert_rows), 3, "у EXPERT-SK должно быть >=3 записей с разными датами"
        )
        dates = {e.date for e in expert_rows}
        self.assertEqual(len(dates), len(expert_rows), "все даты у EXPERT-SK уникальные")

    def test_raises_when_content_column_missing(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Контрагент"
        ws["B1"] = "Дата"
        with self.assertRaises(ValueError):
            parse_workbook_via_temp(wb)


class TestWriteResult(unittest.TestCase):
    def test_combined_fill_used_when_row_has_both(self) -> None:
        wb = _make_minimal_workbook()
        corrections = [
            Correction(row_idx=2, content_col=2, original_content="original", new_content="fixed")
        ]
        warnings = [WarningCell(row_idx=2, content_col=2)]

        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir) / "input.xlsx"
            wb.save(tmp)
            out = write_result(wb, corrections, warnings, tmp, header_row=1, output_dir=Path(tmpdir))
            self.assertTrue(out.exists())

            from openpyxl import load_workbook

            written = load_workbook(out)
            ws = written.active
            self.assertEqual(ws.cell(row=1, column=3).value, "Исправленное содержание")
            self.assertEqual(ws.cell(row=2, column=2).value, "original", "оригинал не должен меняться")
            self.assertEqual(ws.cell(row=2, column=3).value, "fixed")
            self.assertEqual(
                ws.cell(row=2, column=2).fill.start_color.rgb, COMBINED_FILL.start_color.rgb
            )
            self.assertEqual(
                ws.cell(row=2, column=3).fill.start_color.rgb, COMBINED_FILL.start_color.rgb
            )

    def test_separate_fills_for_separate_rows(self) -> None:
        wb = _make_minimal_workbook()
        corrections = [
            Correction(row_idx=2, content_col=2, original_content="original", new_content="fixed")
        ]
        warnings = [WarningCell(row_idx=3, content_col=2)]

        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir) / "input.xlsx"
            wb.save(tmp)
            out = write_result(wb, corrections, warnings, tmp, header_row=1, output_dir=Path(tmpdir))

            from openpyxl import load_workbook

            written = load_workbook(out)
            ws = written.active
            self.assertEqual(ws.cell(row=2, column=2).value, "original")
            self.assertEqual(ws.cell(row=2, column=3).value, "fixed")
            self.assertEqual(
                ws.cell(row=2, column=2).fill.start_color.rgb, HIGHLIGHT_FILL.start_color.rgb
            )
            self.assertEqual(
                ws.cell(row=2, column=3).fill.start_color.rgb, HIGHLIGHT_FILL.start_color.rgb
            )
            self.assertEqual(ws.cell(row=3, column=2).value, "(чей ПК)")
            self.assertIsNone(ws.cell(row=3, column=3).value, "новая колонка пустая для warning")
            self.assertEqual(
                ws.cell(row=3, column=2).fill.start_color.rgb, WARNING_FILL.start_color.rgb
            )
            self.assertEqual(
                ws.cell(row=3, column=3).fill.start_color.rgb, WARNING_FILL.start_color.rgb
            )


    def test_corrected_column_uses_rich_text_with_red_runs(self) -> None:
        wb = _make_minimal_workbook()
        corrections = [
            Correction(
                row_idx=2, content_col=2,
                original_content="обновлене иб", new_content="обновление ИБ",
            )
        ]

        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir) / "input.xlsx"
            wb.save(tmp)
            out = write_result(wb, corrections, [], tmp, header_row=1, output_dir=Path(tmpdir))

            from openpyxl import load_workbook
            from openpyxl.cell.rich_text import CellRichText, TextBlock

            written = load_workbook(out, rich_text=True)
            ws = written.active
            value = ws.cell(row=2, column=3).value
            self.assertIsInstance(value, CellRichText)
            red_blocks = [
                p for p in value
                if isinstance(p, TextBlock) and p.font.color and p.font.color.rgb == "FFCC0000"
            ]
            self.assertTrue(red_blocks, "должны быть красные TextBlock'и для изменённых фрагментов")
            red_text = "".join(b.text for b in red_blocks)
            self.assertIn("обновление", red_text, "слово целиком должно быть красным при пословном diff")
            self.assertIn("ИБ", red_text)

    def test_corrected_cell_wraps_text_and_keeps_source_alignment(self) -> None:
        wb = _make_minimal_workbook()
        ws = wb.active
        ws.column_dimensions["B"].width = 42
        ws["B2"].alignment = Alignment(horizontal="right", vertical="top", indent=1)
        corrections = [
            Correction(row_idx=2, content_col=2, original_content="original", new_content="fixed")
        ]

        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir) / "input.xlsx"
            wb.save(tmp)
            out = write_result(wb, corrections, [], tmp, header_row=1, output_dir=Path(tmpdir))

            from openpyxl import load_workbook

            written = load_workbook(out)
            written_ws = written.active
            corrected = written_ws.cell(row=2, column=3)
            original = written_ws.cell(row=2, column=2)
            self.assertTrue(corrected.alignment.wrap_text)
            self.assertEqual(corrected.alignment.horizontal, "right")
            self.assertEqual(corrected.alignment.vertical, "top")
            self.assertEqual(corrected.alignment.indent, 1)
            self.assertIsNone(original.alignment.wrap_text)
            self.assertEqual(written_ws.column_dimensions["C"].width, 42)


def _make_minimal_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Контрагент"
    ws["B1"] = "Содержание"
    ws["A2"] = "ООО Тест"
    ws["B2"] = "original"
    ws["A3"] = "Иванов ИП"
    ws["B3"] = "(чей ПК)"
    return wb


def parse_workbook_via_temp(wb: Workbook):
    import tempfile

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir) / "input.xlsx"
        wb.save(tmp)
        return parse(tmp)


if __name__ == "__main__":
    unittest.main()
