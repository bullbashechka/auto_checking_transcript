"""Локальный e2e-тест полного конвейера через scripts.check_file."""
from __future__ import annotations

import asyncio
import importlib.util
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

HAS_OPENAI = importlib.util.find_spec("openai") is not None

if HAS_OPENAI:
    from scripts import check_file
    from src.config import Settings
    from src.llm import LLMClient


def _stub_settings() -> "Settings":
    return Settings(
        telegram_token="test-token",
        openai_api_key="test-key",
        allowed_user_ids=frozenset(),
        allow_any=True,
        openai_model="gpt-5.6-luna",
        openai_reasoning_effort="low",
        llm_concurrency=1,
    )


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Время нач"
    sheet["D1"] = "Содержание"
    sheet["A2"] = "ООО Тест"

    rows = [
        "Работа  по извещениям",
        "Обязательство по Извещениям",
        "Учет в с/х. Опечатка",
    ]
    row = 3
    for index, content in enumerate(rows):
        sheet.cell(row=row, column=1).value = "01.01.2026"
        row += 1
        sheet.cell(row=row, column=1).value = f"{9 + index:02d}:00:00"
        sheet.cell(row=row, column=4).value = content
        row += 1

    workbook.save(path)


@unittest.skipUnless(HAS_OPENAI, "openai not installed — run `pip install -r requirements.txt`")
class TestCheckFileE2E(unittest.TestCase):
    def _run(self, coroutine):
        loop = asyncio.new_event_loop()
        try:
            return loop.run_until_complete(coroutine)
        finally:
            loop.close()

    def test_full_local_pipeline_without_network(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "input.xlsx"
            _make_workbook(input_path)

            async def fake_call(_client: LLMClient, user_message: str) -> str:
                payload = json.loads(user_message)
                results = []
                for item in payload:
                    if item["id"] == 2:
                        results.append(
                            {
                                "id": 2,
                                "Исправленное_Содержание": "Учёт в СХ. Исправлена",
                                "Изменения": ["исправлена опечатка"],
                                "Предупреждение": None,
                            }
                        )
                    else:
                        results.append(
                            {
                                "id": item["id"],
                                "Исправленное_Содержание": None,
                                "Изменения": [],
                                "Предупреждение": None,
                            }
                        )
                return json.dumps({"results": results}, ensure_ascii=False)

            with (
                patch.object(check_file, "load_settings", return_value=_stub_settings()),
                patch.object(LLMClient, "_call_with_retries", new=fake_call),
                patch.object(sys, "argv", ["check_file", str(input_path)]),
            ):
                output = io.StringIO()
                with redirect_stdout(output):
                    self._run(check_file.main())

            output_files = list(Path(tmpdir).glob("input_исправленное_*.xlsx"))
            self.assertEqual(len(output_files), 1)
            result_path = output_files[0]
            self.assertIn("Файл с правками:", output.getvalue())

            original = load_workbook(input_path).active
            self.assertEqual(original["D4"].value, "Работа  по извещениям")
            self.assertEqual(original["D6"].value, "Обязательство по Извещениям")

            result = load_workbook(result_path).active
            self.assertEqual(result["E4"].value, "Работа по Извещениям.")
            self.assertEqual(result["E6"].value, "Обязательство по извещениям.")
            self.assertEqual(result["E8"].value, "Учет в с/х. Исправлена.")
            for row in (4, 6, 8):
                self.assertTrue(result.cell(row=row, column=5).alignment.wrap_text)


if __name__ == "__main__":
    unittest.main()
