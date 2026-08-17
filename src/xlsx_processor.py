from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass
from datetime import date
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

DATE_RE = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")


HIGHLIGHT_FILL = PatternFill(start_color="FFFFF2A8", end_color="FFFFF2A8", fill_type="solid")
WARNING_FILL = PatternFill(start_color="FFFFC7C7", end_color="FFFFC7C7", fill_type="solid")
COMBINED_FILL = PatternFill(start_color="FFFFB14D", end_color="FFFFB14D", fill_type="solid")

_DIFF_FONT = InlineFont(color="FFCC0000")
_TOKEN_RE = re.compile(r"\S+|\s+")


def _build_diff_rich_text(original: str, corrected: str) -> CellRichText | str:
    """Возвращает CellRichText с красным шрифтом на изменённых/вставленных словах.
    Diff пословный: токен — это «непробельный» кусок (слово с прилипшей пунктуацией)
    или «пробельный» кусок. Слово целиком становится красным, если оно поменялось."""
    a_tokens = _TOKEN_RE.findall(original)
    b_tokens = _TOKEN_RE.findall(corrected)
    matcher = SequenceMatcher(a=a_tokens, b=b_tokens, autojunk=False)
    parts: list[str | TextBlock] = []
    for op, _i1, _i2, j1, j2 in matcher.get_opcodes():
        if j1 == j2:
            continue
        chunk = "".join(b_tokens[j1:j2])
        if op == "equal":
            parts.append(chunk)
        else:
            parts.append(TextBlock(_DIFF_FONT, chunk))
    if not parts:
        return ""
    if len(parts) == 1 and isinstance(parts[0], str):
        return parts[0]
    return CellRichText(parts)


@dataclass(frozen=True)
class WorkEntry:
    contractor: str
    date: str
    time: str
    content: str
    row_idx: int  # 1-based row in the sheet
    content_col: int  # 1-based column index of «Содержание»


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _find_content_column(ws) -> tuple[int, int]:
    """Returns (header_row, content_col) — 1-based indices."""
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        for col_idx, value in enumerate(row, 1):
            if _cell_text(value).lower() == "содержание":
                return row_idx, col_idx
    raise ValueError("Не найдена колонка «Содержание» в первых 10 строках листа")


def parse(path: Path | str) -> tuple[Workbook, list[WorkEntry], int]:
    """Open the workbook, locate the «Содержание» column, walk the three-line layout.

    Каждая работа в файле — это три подряд идущие строки:
        контрагент (текст в колонке A)
        дата       (dd.mm.yyyy в колонке A)
        работа     (время в колонке A, содержание в колонке «Содержание»)
    У одного контрагента может быть несколько пар «дата + работа» подряд.

    Returns (workbook, entries, header_row).
    """
    wb = load_workbook(path)
    ws = wb.active

    header_row, content_col = _find_content_column(ws)

    entries: list[WorkEntry] = []
    current_contractor = ""
    current_date = ""

    for row_idx in range(header_row + 1, ws.max_row + 1):
        col_a = _cell_text(ws.cell(row=row_idx, column=1).value)
        content_cell = ws.cell(row=row_idx, column=content_col)
        raw_value = content_cell.value
        if isinstance(raw_value, str) and raw_value != raw_value.rstrip():
            content_cell.value = raw_value.rstrip()
        content = _cell_text(raw_value)

        if not col_a and not content:
            continue

        if content:
            entries.append(
                WorkEntry(
                    contractor=current_contractor,
                    date=current_date,
                    time=col_a,
                    content=content,
                    row_idx=row_idx,
                    content_col=content_col,
                )
            )
            continue

        if DATE_RE.match(col_a):
            current_date = col_a
        else:
            current_contractor = col_a
            current_date = ""

    return wb, entries, header_row


@dataclass(frozen=True)
class Correction:
    row_idx: int
    content_col: int
    original_content: str
    new_content: str


@dataclass(frozen=True)
class WarningCell:
    row_idx: int
    content_col: int


def write_result(
    wb: Workbook,
    corrections: list[Correction],
    warnings: list[WarningCell],
    input_path: Path,
    header_row: int,
    output_dir: Path | None = None,
) -> Path:
    """Write corrections to a new rightmost column, highlight original + new cells.

    Оригинальная ячейка «Содержание» не перезаписывается. Исправленный текст
    пишется в новую колонку справа от всех существующих; обе ячейки
    подсвечиваются. Если на строку пришло и исправление, и warning —
    используется COMBINED_FILL.
    """
    ws = wb.active

    new_col = ws.max_column + 1
    ws.cell(row=header_row, column=new_col).value = "Исправленное содержание"

    if corrections:
        source_col_letter = get_column_letter(corrections[0].content_col)
        new_col_letter = get_column_letter(new_col)
        source_width = ws.column_dimensions[source_col_letter].width
        if source_width is not None:
            ws.column_dimensions[new_col_letter].width = source_width

    correction_rows = {c.row_idx for c in corrections}
    warning_rows = {w.row_idx for w in warnings}
    combined_rows = correction_rows & warning_rows

    for corr in corrections:
        fill = COMBINED_FILL if corr.row_idx in combined_rows else HIGHLIGHT_FILL
        new_cell = ws.cell(row=corr.row_idx, column=new_col)
        new_cell.value = _build_diff_rich_text(corr.original_content, corr.new_content)
        alignment = copy(ws.cell(row=corr.row_idx, column=corr.content_col).alignment)
        alignment.wrap_text = True
        new_cell.alignment = alignment
        new_cell.fill = fill
        ws.cell(row=corr.row_idx, column=corr.content_col).fill = fill

    for warn in warnings:
        if warn.row_idx in combined_rows:
            continue
        ws.cell(row=warn.row_idx, column=warn.content_col).fill = WARNING_FILL
        ws.cell(row=warn.row_idx, column=new_col).fill = WARNING_FILL

    stem = input_path.stem
    suffix = date.today().strftime("%Y-%m-%d")
    out_name = f"{stem}_исправленное_{suffix}.xlsx"
    out_dir = output_dir or input_path.parent
    out_path = out_dir / out_name

    wb.save(out_path)
    return out_path
